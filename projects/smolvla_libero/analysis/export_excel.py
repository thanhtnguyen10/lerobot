# Export results to Excel; every summary number is a formula over per-episode sheets (success = 1/0).
#   --run <eval_combined.py --out>  ->  1_vanilla_libero.xlsx, 2_libero_plus_robustness.xlsx
#   --history <runs.json>           ->  3_vanilla_run_history.xlsx (earlier evaluations)
# Usage (openpyxl is not in LeRobot's lock, so it is added for this command only):
#   uv run --no-project --with openpyxl python projects/smolvla_libero/analysis/export_excel.py \
#     --run outputs/eval/robustness --history outputs/eval/archive/runs.json --out outputs/eval/excel

import argparse
import json
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.defined_name import DefinedName

SUITES = {"libero_spatial": "Spatial", "libero_object": "Object", "libero_goal": "Goal", "libero_10": "Long"}
PERTURBATIONS = {"camera": "Camera Viewpoints", "robot": "Robot Initial States", "light": "Light Conditions",
                 "background": "Background Textures", "noise": "Sensor Noise", "language": "Language Instructions",
                 "layout": "Objects Layout"}
IDENTICAL = {"camera": "task, initial state, seed, instruction; only the camera view changes",
             "robot": "task, objects' initial state, seed, instruction; the robot's start pose changes",
             "light": "task, initial state, seed, instruction; only the lighting changes",
             "background": "task, initial state, seed, instruction; only scene textures change",
             "noise": "task, initial state, seed, instruction; only the image noise changes",
             "language": "task, initial state, seed; the instruction is reworded",
             "layout": "task, seed, instruction; the scene layout changes (own initial state)"}
PAPER = {"Spatial": 0.90, "Object": 0.96, "Goal": 0.92, "Long": 0.71, "Average": 0.873}
PAPER_SOURCE = "Shukor et al. (2025), SmolVLA, Table 2: SmolVLA (0.45B) on LIBERO."
KNOWN_ISSUES = {("Spatial", 5): "reset broken under MuJoCo >= 3.4 (LeRobot #4390)"}

# Style
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
INPUT = "0000FF"                                           # blue = hardcoded reference / editable input
PCT, PTS, INT = "0.0%", "+0.0;-0.0;0.0", "0"


# ---------------------------------------------------------------- formatting helpers

def font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = font(bold=True, size=13)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = font(italic=True, color="595959")


def header(ws, row, labels, col=1):
    for j, label in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=label)
        c.font = font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = Border(bottom=THIN)


def put(ws, row, col, value, fmt=None, bold=False, color=None, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font(bold=bold, color=color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def widths(ws, values):
    for j, w in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def name_range(wb, name, sheet, col, first, last):
    wb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!${col}${first}:${col}${last}")


def scale(ws, ref, reverse=False):
    # white -> blue color scale (reverse: blue = smallest)
    lo, hi = ("FFFFFF", "5B9BD5") if not reverse else ("5B9BD5", "FFFFFF")
    ws.conditional_formatting.add(ref, ColorScaleRule(start_type="min", start_color=lo, end_type="max", end_color=hi))


def readme(ws, heading, lines):
    title(ws, heading)
    for i, line in enumerate(lines, 3):
        if isinstance(line, tuple):
            put(ws, i, 1, line[0], bold=True)
            put(ws, i, 2, line[1])
        else:
            put(ws, i, 1, line)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120


def link(cell, path: str, xlsx_dir: Path, repo: Path):
    if path:
        cell.hyperlink = os.path.relpath(repo / path, xlsx_dir)
        cell.font = font(color="0563C1", underline="single")


# ---------------------------------------------------------------- data loading

def load_vanilla(run: Path, tasks: dict) -> list[dict]:
    rows = []
    for suite, short in SUITES.items():
        info = json.loads((run / "vanilla" / suite / "eval_info.json").read_text())
        for t in info["per_task"]:
            videos = t["metrics"].get("video_paths") or []
            for k, s in enumerate(t["metrics"]["successes"]):
                rows.append({"suite": short, "task": t["task_id"], "instruction": tasks[short][t["task_id"]],
                             "episode": k, "seed": 1000 + k, "success": int(bool(s)),
                             "video": videos[k] if k < len(videos) else ""})
    return rows


def load_plus(run: Path) -> list[dict]:
    videos = {}
    for pert in PERTURBATIONS:
        for suite in SUITES:
            info = json.loads((run / "plus" / pert / suite / "eval_info.json").read_text())
            for t in info["per_task"]:
                v = t["metrics"].get("video_paths") or []
                videos[(pert, suite, t["task_id"])] = v[0] if v else ""
    rows = []
    for line in (run / "plus_episodes.jsonl").read_text().splitlines():
        e = json.loads(line)
        rows.append({"suite": SUITES[e["suite"]], "task": e["vanilla_task_id"], "instruction": e["vanilla_instruction"],
                     "perturbation": PERTURBATIONS[e["perturbation"]], "slot": e["slot"], "seed": 1000 + e["slot"],
                     "level": e["level"], "name": e["name"], "task_index": e["task_index"],
                     "success": int(bool(e["success"])),
                     "video": videos.get((e["perturbation"], e["suite"], e["task_index"]), "")})
    sorder = {s: i for i, s in enumerate(SUITES.values())}
    porder = {p: i for i, p in enumerate(PERTURBATIONS.values())}
    rows.sort(key=lambda r: (porder[r["perturbation"]], sorder[r["suite"]], r["task"], r["slot"]))
    return rows


# ---------------------------------------------------------------- shared sheet

def vanilla_episode_sheet(wb, rows, xlsx_dir, repo):
    # 'Vanilla episodes' + named ranges V_SUITE, V_TASK, V_EP, V_SUCC
    ws = wb.create_sheet("Vanilla episodes")
    header(ws, 1, ["Suite", "Task", "Instruction", "Episode k", "Initial state", "Seed", "Success (1/0)", "Video"])
    for i, r in enumerate(rows, 2):
        for j, v in enumerate([r["suite"], r["task"], r["instruction"], r["episode"], r["episode"], r["seed"],
                               r["success"]], 1):
            put(ws, i, j, v)
        link(put(ws, i, 8, r["video"]), r["video"], xlsx_dir, repo)
    last = len(rows) + 1
    for name, col in [("SUITE", "A"), ("TASK", "B"), ("EP", "D"), ("SUCC", "G")]:
        name_range(wb, f"V_{name}", ws.title, col, 2, last)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{last}"
    widths(ws, [10, 6, 70, 10, 12, 8, 13, 60])


# ---------------------------------------------------------------- workbook 1: vanilla LIBERO

def build_vanilla(out: Path, run: Path, repo: Path, vanilla: list[dict], tasks: dict, config: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    v = config.get("versions", {}).get("vanilla", {})
    readme(ws, "Vanilla LIBERO: SmolVLA evaluation (SmolVLA paper protocol)", [
        ("Policy", config["args"]["policy"]),
        ("Protocol", "4 suites x 10 tasks x 10 episodes = 400. Episode k of a task starts from LIBERO's stored initial "
                     "state k with seed 1000+k; new policy call after every executed action; 10 flow steps; 256x256."),
        ("Evaluation", "deterministic GPU kernels, batch size 1, policy noise reseeded every episode "
                       "(eval_vanilla/deterministic_eval.py); a rerun reproduces every episode"),
        ("Versions", ", ".join(f"{k} {val}" for k, val in v.items() if k != "libero_path")),
        ("Source data", os.path.relpath(run / "vanilla", repo)),
        ("Paper numbers", PAPER_SOURCE),
        "",
        ("Sheets", ""),
        ("By suite", "success rate per suite vs the SmolVLA paper"),
        ("By task", "success rate of each of the 40 tasks"),
        ("Vanilla episodes", "one row per episode (raw data)"),
        "",
        ("Formats", "rates are fractions shown as %; Δ in percentage points; blue = hardcoded reference"),
    ])
    vanilla_episode_sheet(wb, vanilla, out, repo)

    # By suite: A Suite | B Episodes | C Successes | D Success % | E Paper | F Δ vs paper
    ws = wb.create_sheet("By suite", 1)
    title(ws, "Vanilla LIBERO success rate by suite", "Δ vs paper in percentage points.")
    header(ws, 4, ["Suite", "Episodes", "Successes", "Success %", "SmolVLA paper", "Δ vs paper (pts)"])
    names = list(SUITES.values())
    for i, s in enumerate(names + ["Average"], 5):
        total = s == "Average"
        fill = TOTAL_FILL if total else None
        put(ws, i, 1, s, bold=total, fill=fill)
        put(ws, i, 2, "=SUM(B5:B8)" if total else f"=COUNTIFS(V_SUITE,A{i})", INT, fill=fill)
        put(ws, i, 3, "=SUM(C5:C8)" if total else f"=SUMIFS(V_SUCC,V_SUITE,A{i})", INT, fill=fill)
        put(ws, i, 4, f"=IF(B{i}=0,\"\",C{i}/B{i})", PCT, bold=total, fill=fill)
        put(ws, i, 5, PAPER[s], PCT, color=INPUT, fill=fill).comment = Comment(PAPER_SOURCE, "export_excel")
        put(ws, i, 6, f"=(D{i}-E{i})*100", PTS, fill=fill)
    put(ws, 11, 1, "Average = pooled over all 400 episodes (= mean of the 4 suites). Paper average = mean of its "
                   "4 suite numbers.")
    widths(ws, [12, 10, 11, 11, 14, 16])
    ws.freeze_panes = "A5"

    # By task: A Suite | B Task | C Instruction | D Episodes | E Successes | F Success %
    ws = wb.create_sheet("By task", 2)
    title(ws, "Vanilla LIBERO success rate by task")
    header(ws, 3, ["Suite", "Task", "Instruction", "Episodes", "Successes", "Success %"])
    i = 4
    for s in names:
        for tid, instr in sorted(tasks[s].items()):
            put(ws, i, 1, s)
            put(ws, i, 2, tid)
            put(ws, i, 3, instr)
            put(ws, i, 4, f"=COUNTIFS(V_SUITE,A{i},V_TASK,B{i})", INT)
            put(ws, i, 5, f"=SUMIFS(V_SUCC,V_SUITE,A{i},V_TASK,B{i})", INT)
            put(ws, i, 6, f"=IF(D{i}=0,\"\",E{i}/D{i})", PCT)
            if (s, tid) in KNOWN_ISSUES:
                ws.cell(row=i, column=6).comment = Comment(KNOWN_ISSUES[(s, tid)], "export_excel")
            i += 1
    scale(ws, f"F4:F{i - 1}")
    ws.auto_filter.ref = f"A3:F{i - 1}"
    widths(ws, [10, 6, 80, 10, 11, 11])
    ws.freeze_panes = "A4"

    wb.calculation.fullCalcOnLoad = True
    path = out / "1_vanilla_libero.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------- workbook 2: LIBERO-Plus robustness

def build_robustness(out: Path, run: Path, repo: Path, vanilla, plus, tasks, config, summary):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    readme(ws, "Robustness: vanilla LIBERO vs LIBERO-Plus (paired episodes)", [
        ("Policy", config["args"]["policy"]),
        ("Design", "LIBERO-Plus episode k of a task uses the same task, initial state k, seed 1000+k and instruction "
                   "as vanilla episode k (where the perturbation allows), so every LIBERO-Plus episode has a vanilla "
                   "twin. Each distinct LIBERO-Plus variant is used once, up to 10 per (task, perturbation)."),
        ("Vanilla", "vanilla success on the twin episodes of the perturbed ones"),
        ("OOD", "success with the LIBERO-Plus perturbation"),
        ("Δ", "OOD − Vanilla, in percentage points"),
        ("Source data", os.path.relpath(run, repo)),
        "",
        ("Sheets", ""),
        ("Robustness gap", "Vanilla / OOD / Δ per perturbation, all suites (main table)"),
        ("Gap by suite", "the same per suite, plus suite x perturbation matrices"),
        ("Task x perturbation", "success of each task, vanilla and per perturbation; variant counts"),
        ("Difficulty level", "OOD, vanilla and Δ by LIBERO-Plus difficulty level (L1 easiest)"),
        ("Task ranking", "candidate tasks for rapid experiments (editable criteria)"),
        ("Short pairs", "(task, perturbation) pairs with fewer than 10 LIBERO-Plus variants"),
        ("LIBERO-Plus episodes", "one row per LIBERO-Plus episode (raw data)"),
        ("Vanilla episodes", "one row per vanilla episode (raw data)"),
        "",
        ("Formats", "rates are fractions shown as %; Δ in percentage points; blue = editable input"),
    ])
    vanilla_episode_sheet(wb, vanilla, out, repo)

    # LIBERO-Plus episodes + named ranges P_*; K = twin success looked up from 'Vanilla episodes'
    ws = wb.create_sheet("LIBERO-Plus episodes")
    header(ws, 1, ["Suite", "Task", "Instruction (vanilla)", "Perturbation", "Slot k", "Seed", "Level",
                   "LIBERO-Plus variant", "LIBERO-Plus task index", "Success (1/0)", "Vanilla twin success (1/0)",
                   "Video"])
    ws["K1"].comment = Comment("From 'Vanilla episodes': same suite, task and episode k = slot k.", "export_excel")
    for i, r in enumerate(plus, 2):
        for j, val in enumerate([r["suite"], r["task"], r["instruction"], r["perturbation"], r["slot"], r["seed"],
                                 r["level"], r["name"], r["task_index"], r["success"]], 1):
            put(ws, i, j, val)
        put(ws, i, 11, f"=SUMIFS(V_SUCC,V_SUITE,A{i},V_TASK,B{i},V_EP,E{i})")
        link(put(ws, i, 12, r["video"]), r["video"], out, repo)
    last = len(plus) + 1
    for name, col in [("SUITE", "A"), ("TASK", "B"), ("PERT", "D"), ("LEVEL", "G"), ("SUCC", "J"), ("TWIN", "K")]:
        name_range(wb, f"P_{name}", ws.title, col, 2, last)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{last}"
    widths(ws, [10, 6, 50, 20, 7, 7, 7, 60, 12, 12, 14, 50])

    # Robustness gap: A Perturbation | B Episodes | C Vanilla succ | D Vanilla % | E OOD succ | F OOD % | G Δ | H Identical
    ws = wb.create_sheet("Robustness gap", 1)
    title(ws, "Robustness gap: Vanilla vs OOD success per perturbation (all suites)",
          "Vanilla = success without perturbation on the same episodes; Δ in percentage points; sorted by Δ.")
    header(ws, 4, ["Perturbation", "Episodes", "Vanilla successes", "Vanilla Success", "OOD successes", "OOD Success",
                   "Δ Success (pts)", "Identical to the vanilla episode"])
    put(ws, 5, 1, "Vanilla (all 400 episodes)", bold=True)
    put(ws, 5, 2, "=COUNT(V_SUCC)", INT)
    put(ws, 5, 5, "=SUM(V_SUCC)", INT)
    put(ws, 5, 6, "=E5/B5", PCT, bold=True)
    first = 6
    perts = sorted(PERTURBATIONS, key=lambda p: summary["overall"][p]["paired"]["delta"])
    for i, p in enumerate(perts, first):
        put(ws, i, 1, PERTURBATIONS[p])
        put(ws, i, 2, f"=COUNTIFS(P_PERT,A{i})", INT)
        put(ws, i, 3, f"=SUMIFS(P_TWIN,P_PERT,A{i})", INT)
        put(ws, i, 4, f"=C{i}/B{i}", PCT)
        put(ws, i, 5, f"=SUMIFS(P_SUCC,P_PERT,A{i})", INT)
        put(ws, i, 6, f"=E{i}/B{i}", PCT, bold=True)
        put(ws, i, 7, f"=(F{i}-D{i})*100", PTS, bold=True)
        put(ws, i, 8, IDENTICAL[p])
    t = first + len(perts)
    put(ws, t, 1, "All 7 perturbations", bold=True, fill=TOTAL_FILL)
    for col in "BCE":
        put(ws, t, column_index_from_string(col), f"=SUM({col}{first}:{col}{t - 1})", INT, fill=TOTAL_FILL)
    put(ws, t, 4, f"=C{t}/B{t}", PCT, fill=TOTAL_FILL)
    put(ws, t, 6, f"=E{t}/B{t}", PCT, bold=True, fill=TOTAL_FILL)
    put(ws, t, 7, f"=(F{t}-D{t})*100", PTS, bold=True, fill=TOTAL_FILL)
    put(ws, t, 8, "", fill=TOTAL_FILL)
    widths(ws, [26, 10, 11, 11, 11, 11, 12, 70])
    ws.freeze_panes = "B5"

    # Gap by suite: long table (A-F) + matrices (K-O)
    ws = wb.create_sheet("Gap by suite", 2)
    title(ws, "Robustness gap per suite", "Vanilla = success without perturbation on the same episodes; Δ in percentage points.")
    header(ws, 4, ["Suite", "Perturbation", "Episodes", "Vanilla Success", "OOD Success", "Δ Success (pts)"])
    i = 5
    for s in SUITES.values():
        for p in PERTURBATIONS.values():
            put(ws, i, 1, s)
            put(ws, i, 2, p)
            put(ws, i, 3, f"=COUNTIFS(P_SUITE,A{i},P_PERT,B{i})", INT)
            put(ws, i, 4, f"=SUMIFS(P_TWIN,P_SUITE,A{i},P_PERT,B{i})/C{i}", PCT)
            put(ws, i, 5, f"=SUMIFS(P_SUCC,P_SUITE,A{i},P_PERT,B{i})/C{i}", PCT)
            put(ws, i, 6, f"=(E{i}-D{i})*100", PTS, bold=True)
            i += 1
    ws.auto_filter.ref = f"A4:F{i - 1}"

    def matrix(top, label, formula, fmt, reverse=False):
        put(ws, top, 11, label, bold=True)
        header(ws, top + 1, ["Perturbation", *SUITES.values()], col=11)
        for a, p in enumerate(PERTURBATIONS.values()):
            r = top + 2 + a
            put(ws, r, 11, p)
            for b in range(len(SUITES)):
                put(ws, r, 12 + b, formula(f"$K{r}", f"{get_column_letter(12 + b)}${top + 1}"), fmt)
        scale(ws, f"L{top + 2}:O{top + 1 + len(PERTURBATIONS)}", reverse)

    matrix(4, "OOD Success", lambda p, s: f"=SUMIFS(P_SUCC,P_SUITE,{s},P_PERT,{p})/COUNTIFS(P_SUITE,{s},P_PERT,{p})",
           PCT)
    matrix(14, "Δ Success (pts)", lambda p, s: f"=(SUMIFS(P_SUCC,P_SUITE,{s},P_PERT,{p})-SUMIFS(P_TWIN,P_SUITE,{s},"
                                               f"P_PERT,{p}))/COUNTIFS(P_SUITE,{s},P_PERT,{p})*100", PTS, reverse=True)
    put(ws, 24, 11, "Vanilla (all episodes)", bold=True)
    header(ws, 25, ["", *SUITES.values()], col=11)
    put(ws, 26, 11, "Success")
    for b in range(len(SUITES)):
        col = get_column_letter(12 + b)
        put(ws, 26, 12 + b, f"=SUMIFS(V_SUCC,V_SUITE,{col}$25)/COUNTIFS(V_SUITE,{col}$25)", PCT)
    widths(ws, [10, 22, 10, 12, 12, 13, 3, 3, 3, 3, 22, 11, 11, 11, 11])
    ws.freeze_panes = "A5"

    # Task x perturbation: D vanilla | E-K OOD per perturbation | L OOD mean | M Δ | N-T variant counts
    ws = wb.create_sheet("Task x perturbation", 3)
    title(ws, "Success by task: vanilla and under each perturbation",
          "OOD mean = unweighted mean of the 7 perturbation rates; Δ = OOD mean − vanilla (pts); "
          "each rate is from up to 10 episodes.")
    pnames = list(PERTURBATIONS.values())
    header(ws, 4, ["Suite", "Task", "Instruction", "Vanilla", *pnames, "OOD mean", "Δ (pts)",
                   *[f"Variants: {p}" for p in pnames]])
    i = 5
    task_rows = {}
    for s in SUITES.values():
        for tid, instr in sorted(tasks[s].items()):
            task_rows[(s, tid)] = i
            put(ws, i, 1, s)
            put(ws, i, 2, tid)
            put(ws, i, 3, instr)
            put(ws, i, 4, f"=AVERAGEIFS(V_SUCC,V_SUITE,$A{i},V_TASK,$B{i})", PCT, bold=True)
            for j in range(len(pnames)):
                col, vcol = get_column_letter(5 + j), get_column_letter(14 + j)
                put(ws, i, 5 + j, f"=IF({vcol}{i}=0,\"\",AVERAGEIFS(P_SUCC,P_SUITE,$A{i},P_TASK,$B{i},P_PERT,"
                                  f"{col}$4))", PCT)
                put(ws, i, 14 + j, f"=COUNTIFS(P_SUITE,$A{i},P_TASK,$B{i},P_PERT,{col}$4)", INT)
            put(ws, i, 12, f"=AVERAGE(E{i}:K{i})", PCT, bold=True)
            put(ws, i, 13, f"=(L{i}-D{i})*100", PTS, bold=True)
            if (s, tid) in KNOWN_ISSUES:
                ws.cell(row=i, column=4).comment = Comment(KNOWN_ISSUES[(s, tid)], "export_excel")
            i += 1
    scale(ws, f"D5:L{i - 1}")
    ws.auto_filter.ref = f"A4:T{i - 1}"
    widths(ws, [9, 6, 60, 9, *[11] * 7, 10, 9, *[11] * 7])
    ws.row_dimensions[4].height = 42
    ws.freeze_panes = "D5"

    # Difficulty level: 4 blocks (OOD, vanilla, Δ, episodes) x levels L1-L5
    ws = wb.create_sheet("Difficulty level", 4)
    title(ws, "By LIBERO-Plus difficulty level (all suites)",
          "LIBERO-Plus labels: how many of 4 reference models solved the variant (L1 easiest, L5 hardest).")
    levels = [1, 2, 3, 4, 5]
    blocks = [("OOD Success", lambda p, lv: f"=IFERROR(AVERAGEIFS(P_SUCC,P_PERT,{p},P_LEVEL,{lv}),\"\")", PCT),
              ("Vanilla Success (same episodes)",
               lambda p, lv: f"=IFERROR(AVERAGEIFS(P_TWIN,P_PERT,{p},P_LEVEL,{lv}),\"\")", PCT),
              ("Δ Success (pts)", lambda p, lv: f"=IFERROR((AVERAGEIFS(P_SUCC,P_PERT,{p},P_LEVEL,{lv})-AVERAGEIFS("
                                                f"P_TWIN,P_PERT,{p},P_LEVEL,{lv}))*100,\"\")", PTS),
              ("Episodes", lambda p, lv: f"=COUNTIFS(P_PERT,{p},P_LEVEL,{lv})", INT)]
    top = 4
    for label, formula, fmt in blocks:
        put(ws, top, 1, label, bold=True)
        header(ws, top + 1, ["Perturbation", *levels])
        for b in range(len(levels)):
            ws.cell(row=top + 1, column=2 + b).number_format = '"L"0'
        for a, p in enumerate(PERTURBATIONS.values()):
            r = top + 2 + a
            put(ws, r, 1, p)
            for b in range(len(levels)):
                put(ws, r, 2 + b, formula(f"$A{r}", f"{get_column_letter(2 + b)}${top + 1}"), fmt)
        if fmt == PCT:
            scale(ws, f"B{top + 2}:F{top + 1 + len(PERTURBATIONS)}")
        top += len(PERTURBATIONS) + 3
    widths(ws, [30, 10, 10, 10, 10, 10])

    # Task ranking: criteria in B3:B4 (editable); rows ordered by the ranking at export time
    ws = wb.create_sheet("Task ranking", 5)
    title(ws, "Candidate tasks for rapid experiments",
          "Qualifies if all criteria hold; qualifying tasks ranked by Δ (largest drop = rank 1).")
    put(ws, 3, 1, "Min vanilla success", bold=True)
    put(ws, 3, 2, 0.8, PCT, color=INPUT).comment = Comment("Editable: minimum vanilla success.", "export_excel")
    put(ws, 4, 1, "Required variants", bold=True)
    put(ws, 4, 2, 10, INT, color=INPUT).comment = Comment("Editable: min LIBERO-Plus variants per perturbation.",
                                                          "export_excel")
    put(ws, 5, 1, "Row order is fixed at export; Qualifies and Rank recompute.")
    header(ws, 7, ["Suite", "Task", "Instruction", "Vanilla Success", "OOD mean", "Δ (pts)", "Min variants",
                   "Known issue", "Qualifies", "Rank"])
    key = {v: k for k, v in SUITES.items()}

    def order_key(s, tid):
        r = summary["by_task"][key[s]][str(tid)]
        van_rate = r["vanilla"]["pc_success"]
        ood = sum(r[p]["pc_success"] for p in PERTURBATIONS) / len(PERTURBATIONS)
        ok = van_rate >= 80 and min(r[p]["n"] for p in PERTURBATIONS) >= 10 and (s, tid) not in KNOWN_ISSUES
        return (not ok, round(ood - van_rate, 6))

    order = sorted(task_rows, key=lambda st: order_key(*st))
    end = 7 + len(order)
    src_sheet = "'Task x perturbation'!"
    for i, (s, tid) in enumerate(order, 8):
        src = task_rows[(s, tid)]
        put(ws, i, 1, f"={src_sheet}A{src}")
        put(ws, i, 2, f"={src_sheet}B{src}")
        put(ws, i, 3, f"={src_sheet}C{src}")
        put(ws, i, 4, f"={src_sheet}D{src}", PCT)
        put(ws, i, 5, f"={src_sheet}L{src}", PCT)
        put(ws, i, 6, f"=ROUND({src_sheet}M{src},6)", PTS, bold=True)  # round so equal Δs tie
        put(ws, i, 7, f"=MIN({src_sheet}N{src}:T{src})", INT)
        put(ws, i, 8, KNOWN_ISSUES.get((s, tid), ""))
        put(ws, i, 9, f"=AND(D{i}>=$B$3,G{i}>=$B$4,H{i}=\"\")")
        put(ws, i, 10, f"=IF(I{i},COUNTIFS($I$8:$I${end},TRUE,$F$8:$F${end},\"<\"&F{i})+1,\"\")", INT, bold=True)
    ws.auto_filter.ref = f"A7:J{end}"
    widths(ws, [18, 6, 70, 12, 10, 9, 12, 44, 10, 7])
    ws.freeze_panes = "A8"

    # Short pairs
    ws = wb.create_sheet("Short pairs", 6)
    title(ws, "(Task, perturbation) pairs with fewer than 10 LIBERO-Plus variants",
          "All available variants are used, each once.")
    header(ws, 4, ["Suite", "Task", "Perturbation", "Episodes"])
    counts = {}
    for r in plus:
        k = (r["suite"], r["task"], r["perturbation"])
        counts[k] = counts.get(k, 0) + 1
    i = 5
    for (s, tid, p), n in counts.items():
        if n < 10:
            put(ws, i, 1, s)
            put(ws, i, 2, tid)
            put(ws, i, 3, p)
            put(ws, i, 4, f"=COUNTIFS(P_SUITE,A{i},P_TASK,B{i},P_PERT,C{i})", INT)
            i += 1
    widths(ws, [10, 6, 22, 10])

    wb.move_sheet("Vanilla episodes", offset=len(wb.sheetnames))
    wb.calculation.fullCalcOnLoad = True
    path = out / "2_libero_plus_robustness.xlsx"
    wb.save(path)
    return path


# ---------------------------------------------------------------- workbook 3: run history

def build_history(out: Path, spec_path: Path, tasks: dict):
    spec = json.loads(spec_path.read_text())
    episodes = []
    for run in spec["runs"]:
        for d in run["dirs"]:
            info = json.loads((spec_path.parent / d / "eval_info.json").read_text())
            for t in info["per_task"]:
                for k, s in enumerate(t["metrics"]["successes"]):
                    episodes.append((run["label"], SUITES[t["task_group"]], t["task_id"], k, int(bool(s))))

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    lines = [("About", spec["_about"]), "", ("Checkpoints", "")]
    lines += [(k, v) for k, v in spec["checkpoints"].items()]
    lines += ["", ("Evaluation modes", "")] + [(k, v) for k, v in spec["modes"].items()]
    lines += ["", ("Sheets", ""),
              ("Runs", "every evaluation: checkpoint, MuJoCo version, mode, eval seed"),
              ("By suite", "success per run and suite, with the SmolVLA paper for reference"),
              ("By task", "success per task (columns = runs)"),
              ("Determinism", "episode-by-episode comparison of the determinism checks (Spatial tasks 0-4)"),
              ("Episodes", "one row per episode of every run (raw data)"),
              "", ("Paper numbers", PAPER_SOURCE)]
    readme(ws, "Vanilla LIBERO: all evaluations (history and checks)", lines)

    # Episodes + named ranges H_*
    ws = wb.create_sheet("Episodes")
    header(ws, 1, ["Run", "Suite", "Task", "Episode k", "Success (1/0)"])
    for i, row in enumerate(episodes, 2):
        for j, v in enumerate(row, 1):
            put(ws, i, j, v)
    last = len(episodes) + 1
    for name, col in [("RUN", "A"), ("SUITE", "B"), ("TASK", "C"), ("EP", "D"), ("SUCC", "E")]:
        name_range(wb, f"H_{name}", ws.title, col, 2, last)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{last}"
    widths(ws, [34, 9, 6, 10, 13])

    ws = wb.create_sheet("Runs", 1)
    title(ws, "Evaluations")
    header(ws, 3, ["Run", "Group", "Checkpoint", "MuJoCo", "Mode", "Eval seed", "Episodes", "Note", "Source dir(s)"])
    for i, run in enumerate(spec["runs"], 4):
        for j, v in enumerate([run["label"], run["group"], run["checkpoint"], run["mujoco"], run["mode"],
                               run["eval_seed"]], 1):
            put(ws, i, j, v)
        put(ws, i, 7, f"=COUNTIFS(H_RUN,A{i})", INT)
        put(ws, i, 8, run["note"])
        put(ws, i, 9, ", ".join(run["dirs"]))
    widths(ws, [34, 20, 11, 9, 9, 10, 10, 50, 60])
    ws.freeze_panes = "B4"

    ws = wb.create_sheet("By suite", 2)
    title(ws, "Success by run and suite", "Blank = suite not evaluated. Average = pooled over the run's episodes.")
    names = list(SUITES.values())
    header(ws, 4, ["Run", "Group", *names, "Average", "Episodes"])
    put(ws, 5, 1, "SmolVLA paper", bold=True)
    put(ws, 5, 2, "Reference")
    for b, s in enumerate(names + ["Average"]):
        put(ws, 5, 3 + b, PAPER[s], PCT, color=INPUT).comment = Comment(PAPER_SOURCE, "export_excel")
    for i, run in enumerate(spec["runs"], 6):
        put(ws, i, 1, run["label"], bold=run["group"] == "Reported result")
        put(ws, i, 2, run["group"])
        for b in range(len(names)):
            put(ws, i, 3 + b, f"=IFERROR(AVERAGEIFS(H_SUCC,H_RUN,$A{i},H_SUITE,{get_column_letter(3 + b)}$4),\"\")",
                PCT)
        put(ws, i, 7, f"=IFERROR(AVERAGEIFS(H_SUCC,H_RUN,$A{i}),\"\")", PCT, bold=True)
        put(ws, i, 8, f"=COUNTIFS(H_RUN,$A{i})", INT)
    scale(ws, f"C5:G{5 + len(spec['runs'])}")
    widths(ws, [34, 20, 10, 10, 10, 10, 10, 10])
    ws.freeze_panes = "C5"

    ws = wb.create_sheet("By task", 3)
    title(ws, "Success by task (columns = runs)", "Blank = task not evaluated in that run.")
    header(ws, 3, ["Suite", "Task", "Instruction", *[r["label"] for r in spec["runs"]]])
    ws.row_dimensions[3].height = 60
    i = 4
    for s in names:
        for tid, instr in sorted(tasks[s].items()):
            put(ws, i, 1, s)
            put(ws, i, 2, tid)
            put(ws, i, 3, instr)
            for b in range(len(spec["runs"])):
                col = get_column_letter(4 + b)
                put(ws, i, 4 + b, f"=IFERROR(AVERAGEIFS(H_SUCC,H_RUN,{col}$3,H_SUITE,$A{i},H_TASK,$B{i}),\"\")", PCT)
            i += 1
    scale(ws, f"D4:{get_column_letter(3 + len(spec['runs']))}{i - 1}")
    widths(ws, [9, 6, 60, *[12] * len(spec["runs"])])
    ws.freeze_panes = "D4"

    # Determinism: one column per check run; bottom rows count differing episodes
    det = [r["label"] for r in spec["runs"] if r["group"] == "Determinism check"]
    ws = wb.create_sheet("Determinism", 4)
    title(ws, "Determinism checks: episode-by-episode outcomes (Spatial tasks 0-4)",
          "Each column is one evaluation of the same checkpoint.")
    header(ws, 4, ["Suite", "Task", "Episode k", *det])
    ws.row_dimensions[4].height = 45
    i = 5
    for tid in range(5):
        for k in range(10):
            put(ws, i, 1, "Spatial")
            put(ws, i, 2, tid)
            put(ws, i, 3, k)
            for b in range(len(det)):
                col = get_column_letter(4 + b)
                put(ws, i, 4 + b, f"=SUMIFS(H_SUCC,H_RUN,{col}$4,H_SUITE,$A{i},H_TASK,$B{i},H_EP,$C{i})", INT)
            i += 1
    last = i - 1
    i += 1
    put(ws, i, 1, "Successes", bold=True)
    for b in range(len(det)):
        col = get_column_letter(4 + b)
        put(ws, i, 4 + b, f"=SUM({col}5:{col}{last})", INT, bold=True)
    i += 2
    put(ws, i, 1, "Episodes that differ", bold=True)
    for a, b in [(a, b) for a, b in [(0, 1), (2, 3), (2, 4), (5, 6), (0, 2), (2, 5)] if b < len(det)]:
        i += 1
        ca, cb = get_column_letter(4 + a), get_column_letter(4 + b)
        put(ws, i, 1, f"{det[a]}  vs  {det[b]}")
        put(ws, i, 4, f"=SUMPRODUCT(ABS({ca}5:{ca}{last}-{cb}5:{cb}{last}))", INT, bold=True)
    widths(ws, [44, 6, 10, *[13] * len(det)])
    ws.freeze_panes = "D5"

    wb.move_sheet("Episodes", offset=len(wb.sheetnames))
    wb.calculation.fullCalcOnLoad = True
    path = out / "3_vanilla_run_history.xlsx"
    wb.save(path)
    return path


def main():
    parser = argparse.ArgumentParser(description="Export results to Excel")
    parser.add_argument("--run", type=Path, help="--out dir of an eval_combined.py run")
    parser.add_argument("--history", type=Path, help="runs.json describing earlier evaluations")
    parser.add_argument("--out", type=Path, required=True, help="output dir for the .xlsx files")
    args = parser.parse_args()
    repo = Path(__file__).resolve().parents[3]
    args.out.mkdir(parents=True, exist_ok=True)
    out = args.out.resolve()

    run = (args.run or repo / "outputs/eval/robustness").resolve()  # also the source of task instructions
    tasks = {}
    for line in (run / "plus_episodes.jsonl").read_text().splitlines():
        e = json.loads(line)
        tasks.setdefault(SUITES[e["suite"]], {})[e["vanilla_task_id"]] = e["vanilla_instruction"]

    written = []
    if args.run:
        config = json.loads((run / "run_config.json").read_text())
        summary = json.loads((run / "summary.json").read_text())
        vanilla = load_vanilla(run, tasks)
        plus = load_plus(run)
        written.append(build_vanilla(out, run, repo, vanilla, tasks, config))
        written.append(build_robustness(out, run, repo, vanilla, plus, tasks, config, summary))
    if args.history:
        written.append(build_history(out, args.history.resolve(), tasks))
    for p in written:
        print(f"wrote {p}")


if __name__ == "__main__":
    main()
